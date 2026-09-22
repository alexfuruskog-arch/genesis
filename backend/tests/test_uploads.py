from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient

from app.main import app
from app.services.membership import FREE_SCAN_CAP


client = TestClient(app)


def _auth() -> dict[str, str]:
    signup = client.post(
        "/auth/signup",
        json={"email": "lectures@arzive.org", "password": "secret12"},
    )
    assert signup.status_code == 200, signup.text
    return {"Authorization": f"Bearer {signup.json()['token']}"}


def _pdf_bytes(text: str = "Lecture notes on gravity and orbits.") -> bytes:
    # Minimal PDF with a Helvetica text stream. pypdf can usually extract this.
    body = f"BT /F1 12 Tf 40 100 Td ({text}) Tj ET".encode("latin-1", errors="replace")
    return (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Count 1/Kids[3 0 R]>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>endobj\n"
        b"4 0 obj<</Length " + str(len(body)).encode() + b">>stream\n" + body + b"\nendstream\nendobj\n"
        b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
        b"trailer<</Root 1 0 R>>\n%%EOF\n"
    )


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(text: str) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = text
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_lecture_pdf_lands_in_same_library():
    headers = _auth()
    upload = client.post(
        "/upload-document",
        headers=headers,
        files={"file": ("week-1-gravity.pdf", _pdf_bytes(), "application/pdf")},
        data={"source_kind": "lecture"},
    )
    assert upload.status_code == 200, upload.text
    body = upload.json()
    assert body["source_kind"] == "lecture"
    assert body["kind_label"] == "PDF"
    assert body["notebook_model"] is None
    assert body["storage_backend"] == "local"
    assert body["filename"] == "week-1-gravity.pdf"

    listed = client.get("/documents", headers=headers)
    assert listed.status_code == 200
    docs = listed.json()["documents"]
    assert len(docs) == 1
    assert docs[0]["id"] == body["id"]
    assert docs[0]["source_kind"] == "lecture"

    downloaded = client.get(f"/documents/{body['id']}/file", headers=headers)
    assert downloaded.status_code == 200
    assert downloaded.content == _pdf_bytes()
    assert "application/pdf" in downloaded.headers.get("content-type", "")


def test_docx_and_pptx_extract_and_list():
    headers = _auth()
    docx = client.post(
        "/upload-document",
        headers=headers,
        files={
            "file": (
                "mito.docx",
                _docx_bytes("Mitochondria produce ATP in eukaryotic cells."),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"source_kind": "lecture"},
    )
    assert docx.status_code == 200, docx.text
    assert docx.json()["kind_label"] == "Word"
    assert docx.json()["extraction_mode"] == "extracted"
    assert docx.json()["chunk_count"] >= 1

    pptx = client.post(
        "/upload-document",
        headers=headers,
        files={
            "file": (
                "week2.pptx",
                _pptx_bytes("Newton's second law: F = ma"),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
        data={"source_kind": "lecture"},
    )
    assert pptx.status_code == 200, pptx.text
    assert pptx.json()["kind_label"] == "PowerPoint"
    assert pptx.json()["extraction_mode"] == "extracted"


def test_image_is_stored_without_waiting_on_ocr():
    headers = _auth()
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\nIDATx\x9cc```\x00\x00"
        b"\x00\x02\x00\x01\xe2!\xbc3\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    upload = client.post(
        "/upload-document",
        headers=headers,
        files={"file": ("slide-photo.png", png, "image/png")},
        data={"source_kind": "lecture"},
    )
    assert upload.status_code == 200, upload.text
    assert upload.json()["source_kind"] == "lecture"
    assert upload.json()["kind_label"] == "Image"
    listed = client.get("/documents", headers=headers).json()["documents"]
    assert listed[0]["filename"] == "slide-photo.png"
    file_resp = client.get(f"/documents/{upload.json()['id']}/file", headers=headers)
    assert file_resp.status_code == 200
    assert file_resp.content == png


def test_notebook_scan_still_tags_cornell():
    headers = _auth()
    upload = client.post(
        "/upload-document",
        headers=headers,
        files={"file": ("page.txt", b"Cue column. Main notes. Summary.", "text/plain")},
        data={"notebook_model": "cornell", "source_kind": "scan"},
    )
    assert upload.status_code == 200, upload.text
    assert upload.json()["source_kind"] == "scan"
    assert upload.json()["notebook_model"] == "cornell"


def test_unsupported_type_is_rejected():
    headers = _auth()
    upload = client.post(
        "/upload-document",
        headers=headers,
        files={"file": ("notes.zip", b"PK\x03\x04not-a-real-zip", "application/zip")},
        data={"source_kind": "lecture"},
    )
    assert upload.status_code == 415
    assert "Unsupported file" in upload.json()["detail"]


def test_empty_file_is_rejected():
    headers = _auth()
    upload = client.post(
        "/upload-document",
        headers=headers,
        files={"file": ("empty.pdf", b"", "application/pdf")},
        data={"source_kind": "lecture"},
    )
    assert upload.status_code == 400
    assert "empty" in upload.json()["detail"].lower()


def test_free_library_cap_counts_lecture_uploads():
    headers = _auth()
    for index in range(FREE_SCAN_CAP):
        upload = client.post(
            "/upload-document",
            headers=headers,
            files={"file": (f"lec-{index}.txt", b"Lecture body " + str(index).encode(), "text/plain")},
            data={"source_kind": "lecture"},
        )
        assert upload.status_code == 200, upload.text

    blocked = client.post(
        "/upload-document",
        headers=headers,
        files={"file": ("lec-extra.pdf", _pdf_bytes(), "application/pdf")},
        data={"source_kind": "lecture"},
    )
    assert blocked.status_code == 403
    detail = blocked.json()["detail"]
    assert "5 library items" in detail
    assert "lecture files" in detail

    membership = client.get("/membership/me", headers=headers)
    assert membership.status_code == 200
    snap = membership.json()
    assert snap["scans_used"] == FREE_SCAN_CAP
    assert snap["scans_remaining"] == 0
