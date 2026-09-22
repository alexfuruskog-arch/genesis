from __future__ import annotations

from io import BytesIO

from app.services.filetypes import is_image


def extract_text(filename: str, mime_type: str | None, data: bytes) -> tuple[str, str]:
    """Return (text, mode) where mode is extracted or stubbed."""
    name = filename.lower()
    mime = (mime_type or "").lower()
    if name.endswith((".txt", ".md", ".csv")) or mime.startswith("text/"):
        return data.decode("utf-8", errors="replace"), "extracted"
    if name.endswith(".pdf") or mime == "application/pdf":
        return _extract_pdf(filename, data)
    if name.endswith(".docx") or mime.endswith("wordprocessingml.document"):
        return _extract_docx(filename, data)
    if name.endswith(".pptx") or mime.endswith("presentationml.presentation"):
        return _extract_pptx(filename, data)
    if is_image(filename, mime_type):
        return (
            f"[ARZIVE stored {filename} in your library. Image OCR is not connected yet, "
            "so this photo has no searchable text.]",
            "stubbed",
        )
    return (
        f"[ARZIVE stored {filename} ({mime or 'unknown type'}) without a text parser. "
        "The file is in your library; extraction can be added later.]",
        "stubbed",
    )


def _extract_pdf(filename: str, data: bytes) -> tuple[str, str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(part for part in pages if part)
        if text.strip():
            return text, "extracted"
        return (
            f"[PDF {filename} has no extractable text. OCR can be added later. "
            "The file is stored in your library.]",
            "stubbed",
        )
    except Exception as exc:  # pragma: no cover - defensive parser path
        return f"[PDF extraction failed for {filename}: {exc}]", "stubbed"


def _extract_docx(filename: str, data: bytes) -> tuple[str, str]:
    try:
        from docx import Document

        document = Document(BytesIO(data))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts)
        if text.strip():
            return text, "extracted"
        return (
            f"[Word file {filename} has no extractable text. The file is stored in your library.]",
            "stubbed",
        )
    except Exception as exc:  # pragma: no cover
        return f"[Word extraction failed for {filename}: {exc}]", "stubbed"


def _extract_pptx(filename: str, data: bytes) -> tuple[str, str]:
    try:
        from pptx import Presentation

        deck = Presentation(BytesIO(data))
        parts: list[str] = []
        for index, slide in enumerate(deck.slides, start=1):
            bits: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = (shape.text or "").strip()
                    if value:
                        bits.append(value)
            if bits:
                parts.append(f"Slide {index}\n" + "\n".join(bits))
        text = "\n\n".join(parts)
        if text.strip():
            return text, "extracted"
        return (
            f"[PowerPoint {filename} has no extractable text. The file is stored in your library.]",
            "stubbed",
        )
    except Exception as exc:  # pragma: no cover
        return f"[PowerPoint extraction failed for {filename}: {exc}]", "stubbed"


def chunk_text(text: str, max_chars: int = 800) -> list[str]:
    normalized = " ".join(text.split())
    if not normalized:
        return []
    paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs or [normalized]:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= max_chars:
            current = paragraph
            continue
        for start in range(0, len(paragraph), max_chars):
            piece = paragraph[start : start + max_chars].strip()
            if piece:
                chunks.append(piece)
        current = ""
    if current:
        chunks.append(current)
    return chunks or [normalized[:max_chars]]
